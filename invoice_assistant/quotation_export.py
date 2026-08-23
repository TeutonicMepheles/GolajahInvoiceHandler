from __future__ import annotations

from copy import deepcopy
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from . import AppError


DEFAULT_TEMPLATE = Path(__file__).resolve().parent / "templates" / "quotation_template.docx"
MAX_ITEMS = 100
MAX_CHILDREN = 100
MAX_DESCRIPTION_LENGTH = 1000
PRESERVE_TEMPLATE_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "docProps/core.xml",
    "word/_rels/document.xml.rels",
    "word/settings.xml",
    "word/styles.xml",
}


def _single_line(value, label: str, *, maximum: int, required: bool = False) -> str:
    text = re.sub(r"[\r\n\t]+", " ", str(value or "")).strip()
    if required and not text:
        raise AppError(f"{label}不能为空。", 400, "invalid_quotation_export")
    if len(text) > maximum:
        raise AppError(f"{label}不能超过 {maximum} 个字符。", 400, "invalid_quotation_export")
    return text


def _multiline_text(value, label: str, *, maximum: int) -> str:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = "\n".join(line.rstrip() for line in text.split("\n")).strip()
    if len(text) > maximum:
        raise AppError(f"{label}不能超过 {maximum} 个字符。", 400, "invalid_quotation_export")
    return text


def _integer_money(value, label: str, *, positive: bool = False) -> int:
    if isinstance(value, bool) or value in (None, ""):
        raise AppError(f"{label}格式无效。", 400, "invalid_quotation_export")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise AppError(f"{label}格式无效。", 400, "invalid_quotation_export")
    if not number.is_finite() or number != number.to_integral_value():
        raise AppError(f"{label}必须是整数金额。", 400, "invalid_quotation_export")
    amount = int(number)
    if amount < 0 or (positive and amount <= 0):
        raise AppError(f"{label}必须大于零。", 400, "invalid_quotation_export")
    return amount


def _rate(value, label: str) -> Decimal:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise AppError(f"{label}格式无效。", 400, "invalid_quotation_export")
    if not number.is_finite() or number < 0 or number > 100:
        raise AppError(f"{label}必须在 0% 到 100% 之间。", 400, "invalid_quotation_export")
    return number


def _rate_text(value: Decimal) -> str:
    return format(value.normalize(), "f")


def validate_export_payload(payload: dict) -> dict:
    quotation = payload.get("quotation")
    if not isinstance(quotation, dict):
        raise AppError("缺少报价测算结果。", 400, "invalid_quotation_export")

    quote_date_text = _single_line(payload.get("quote_date") or date.today().isoformat(), "报价日期", maximum=20)
    try:
        quote_date = date.fromisoformat(quote_date_text)
    except ValueError:
        raise AppError("报价日期格式无效。", 400, "invalid_quotation_export")

    title = _single_line(payload.get("title") or "项目开发报价单", "报价单标题", maximum=60, required=True)
    client_name = _single_line(payload.get("client_name"), "需求方名称", maximum=80)
    project_manager = _single_line(payload.get("project_manager"), "项目负责人", maximum=40)
    include_descriptions = payload.get("include_descriptions", False)
    if not isinstance(include_descriptions, bool):
        raise AppError("说明列设置格式无效。", 400, "invalid_quotation_export")

    basic_cost = _integer_money(quotation.get("basic_cost"), "基础开发费用", positive=True)
    other_cost = _integer_money(quotation.get("other_cost"), "其他费用")
    contract_total = _integer_money(quotation.get("contract_total"), "合同总额", positive=True)
    rounding_unit = _integer_money(quotation.get("rounding_unit"), "取整粒度", positive=True)
    if rounding_unit not in (100, 500, 1000):
        raise AppError("取整粒度仅支持 100、500 或 1000 元。", 400, "invalid_quotation_export")

    fees = quotation.get("fees")
    rates = quotation.get("rates")
    if not isinstance(fees, dict) or not isinstance(rates, dict):
        raise AppError("其他费用明细不完整。", 400, "invalid_quotation_export")
    normalized_fees = {
        "vat": _integer_money(fees.get("vat"), "增值税"),
        "surcharge": _integer_money(fees.get("surcharge"), "附加税"),
        "management": _integer_money(fees.get("management"), "管理费"),
    }
    normalized_rates = {
        "vat": _rate(rates.get("vat"), "增值税率"),
        "surcharge": _rate(rates.get("surcharge"), "附加税率"),
        "management": _rate(rates.get("management"), "管理费比例"),
    }

    raw_items = quotation.get("items")
    if not isinstance(raw_items, list) or not raw_items or len(raw_items) > MAX_ITEMS:
        raise AppError("基础开发费用条目数量无效。", 400, "invalid_quotation_export")
    items = []
    for item_index, raw_item in enumerate(raw_items, start=1):
        if not isinstance(raw_item, dict):
            raise AppError("基础开发费用条目格式无效。", 400, "invalid_quotation_export")
        item_amount = _integer_money(raw_item.get("amount"), f"第 {item_index} 个条目金额", positive=True)
        if item_amount % rounding_unit:
            raise AppError("条目金额必须符合所选取整粒度。", 400, "invalid_quotation_export")
        name = _single_line(raw_item.get("name") or f"开发条目 {item_index}", "条目名称", maximum=80, required=True)
        description = _multiline_text(
            raw_item.get("description"),
            f"第 {item_index} 个条目说明",
            maximum=MAX_DESCRIPTION_LENGTH,
        )
        raw_children = raw_item.get("children") or []
        if not isinstance(raw_children, list) or len(raw_children) > MAX_CHILDREN:
            raise AppError("子项目数量无效。", 400, "invalid_quotation_export")
        children = []
        for child_index, raw_child in enumerate(raw_children, start=1):
            if not isinstance(raw_child, dict):
                raise AppError("子项目格式无效。", 400, "invalid_quotation_export")
            child_amount = _integer_money(
                raw_child.get("amount"),
                f"第 {item_index}.{child_index} 个子项目金额",
                positive=True,
            )
            if child_amount % rounding_unit:
                raise AppError("子项目金额必须符合所选取整粒度。", 400, "invalid_quotation_export")
            children.append(
                {
                    "name": _single_line(
                        raw_child.get("name") or f"子项目 {child_index}",
                        "子项目名称",
                        maximum=80,
                        required=True,
                    ),
                    "description": _multiline_text(
                        raw_child.get("description"),
                        f"第 {item_index}.{child_index} 个子项目说明",
                        maximum=MAX_DESCRIPTION_LENGTH,
                    ),
                    "amount": child_amount,
                }
            )
        if children and sum(child["amount"] for child in children) != item_amount:
            raise AppError(f"第 {item_index} 个条目的子项目金额合计不一致。", 400, "invalid_quotation_export")
        items.append(
            {
                "name": name,
                "description": description,
                "amount": item_amount,
                "children": children,
            }
        )

    if sum(item["amount"] for item in items) != basic_cost:
        raise AppError("基础开发费用明细合计不一致。", 400, "invalid_quotation_export")
    if sum(normalized_fees.values()) != other_cost:
        raise AppError("其他费用明细合计不一致。", 400, "invalid_quotation_export")
    if basic_cost + other_cost != contract_total:
        raise AppError("合同总额与费用合计不一致。", 400, "invalid_quotation_export")

    return {
        "title": title,
        "client_name": client_name,
        "project_manager": project_manager,
        "quote_date": quote_date,
        "include_descriptions": include_descriptions,
        "basic_cost": basic_cost,
        "other_cost": other_cost,
        "contract_total": contract_total,
        "rounding_unit": rounding_unit,
        "fees": normalized_fees,
        "rates": normalized_rates,
        "items": items,
    }


def _set_run_font(run, *, chinese: str, latin: str, size: float, bold=None, italic=None) -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), chinese)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _replace_paragraph(
    paragraph,
    text: str,
    *,
    alignment,
    chinese: str = "宋体",
    latin: str = "Times New Roman",
    size: float = 12,
    bold: bool = False,
    italic: bool = False,
) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    run = paragraph.add_run(text)
    _set_run_font(run, chinese=chinese, latin=latin, size=size, bold=bold, italic=italic)
    paragraph.alignment = alignment


def _set_cell_text(cell, text: str, *, bold: bool = False, alignment=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    paragraph = cell.paragraphs[0]
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)
    _replace_paragraph(
        paragraph,
        text,
        alignment=alignment,
        chinese="黑体" if bold else "宋体",
        latin="Arial" if bold else "Times New Roman",
        size=12,
        bold=bold,
    )
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _remove_vertical_merge(cell) -> None:
    properties = cell._tc.get_or_add_tcPr()
    vertical_merge = properties.find(qn("w:vMerge"))
    if vertical_merge is not None:
        properties.remove(vertical_merge)


def _remove_element(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _restore_template_parts(document: BytesIO, template: Path) -> BytesIO:
    document.seek(0)
    restored = BytesIO()
    with ZipFile(document, "r") as generated, ZipFile(template, "r") as source, ZipFile(
        restored, "w", compression=ZIP_DEFLATED
    ) as output:
        source_names = set(source.namelist())
        for info in generated.infolist():
            if info.filename in PRESERVE_TEMPLATE_PARTS and info.filename in source_names:
                payload = source.read(info.filename)
            else:
                payload = generated.read(info.filename)
            output.writestr(info, payload)
    restored.seek(0)
    return restored


def _populate_basic_table(table, data: dict) -> None:
    sample_row_xml = deepcopy(table.rows[2]._tr)
    for row in list(table.rows[2:]):
        _remove_element(row._tr)

    for item in data["items"]:
        children = item["children"] or [
            {
                "name": "",
                "description": item["description"],
                "amount": item["amount"],
            }
        ]
        start_row = len(table.rows)
        for child in children:
            row_xml = deepcopy(sample_row_xml)
            table._tbl.append(row_xml)
            row = table.rows[-1]
            _remove_vertical_merge(row.cells[0])
            _set_cell_text(row.cells[0], "")
            _set_cell_text(row.cells[1], child["name"])
            _set_cell_text(row.cells[2], child["description"], alignment=WD_ALIGN_PARAGRAPH.LEFT)
            _set_cell_text(row.cells[3], str(child["amount"]))
        end_row = len(table.rows) - 1
        if end_row > start_row:
            block_cell = table.cell(start_row, 0).merge(table.cell(end_row, 0))
        else:
            block_cell = table.cell(start_row, 0)
        _set_cell_text(block_cell, item["name"])

    if not data["include_descriptions"]:
        for row_index, row in enumerate(table.rows[1:], start=1):
            project_text = row.cells[1].text
            project_cell = row.cells[1].merge(row.cells[2])
            _set_cell_text(project_cell, project_text, bold=row_index == 1)


def _populate_other_fees_table(table, data: dict) -> None:
    rates = data["rates"]
    fees = data["fees"]
    _set_cell_text(table.cell(1, 0), "税收")
    _set_cell_text(
        table.cell(1, 1),
        f"合同总额{_rate_text(rates['vat'])}%的增值税税额（开具增值税普通发票）",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _set_cell_text(table.cell(1, 2), str(fees["vat"]))
    _set_cell_text(table.cell(2, 0), "税收")
    _set_cell_text(
        table.cell(2, 1),
        f"附加税税额为增值税税额的{_rate_text(rates['surcharge'])}%",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _set_cell_text(table.cell(2, 2), str(fees["surcharge"]))
    _set_cell_text(table.cell(3, 0), "管理费")
    _set_cell_text(
        table.cell(3, 1),
        f"合同总额{_rate_text(rates['management'])}%的学校管理费",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    _set_cell_text(table.cell(3, 2), str(fees["management"]))
    _set_cell_text(table.cell(4, 0), "总计", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    _set_cell_text(table.cell(4, 2), str(data["contract_total"]), bold=True)


def build_quotation_docx(payload: dict, template_path: str | Path | None = None) -> tuple[BytesIO, str]:
    data = validate_export_payload(payload)
    template = Path(template_path or DEFAULT_TEMPLATE).resolve()
    if not template.is_file():
        raise AppError("报价单模板不存在，请重新配置模板文件。", 500, "quotation_template_missing")

    doc = Document(template)
    if len(doc.paragraphs) < 21 or len(doc.tables) < 3:
        raise AppError("报价单模板结构不完整。", 500, "quotation_template_invalid")

    original_paragraphs = list(doc.paragraphs)
    _replace_paragraph(
        original_paragraphs[0],
        data["title"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        chinese="黑体",
        latin="Arial",
        size=14,
        bold=True,
    )
    introduction = (
        f"根据{data['client_name']}的需求，特制定项目开发价格方案如下："
        if data["client_name"]
        else "根据项目需求，特制定项目开发价格方案如下："
    )
    _replace_paragraph(
        original_paragraphs[1],
        introduction,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        chinese="宋体",
        latin="Times New Roman",
        size=12,
    )
    original_paragraphs[1].paragraph_format.line_spacing = 1.5

    _populate_basic_table(doc.tables[0], data)
    _populate_other_fees_table(doc.tables[1], data)
    _remove_element(doc.tables[2]._element)

    _replace_paragraph(
        original_paragraphs[3],
        "备注：以上价格方案仅供参考，最终需以合同订立情况为准。",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        chinese="宋体",
        latin="Times New Roman",
        size=12,
        italic=True,
    )
    _replace_paragraph(
        original_paragraphs[13],
        f"项目负责人：{data['project_manager']}",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        chinese="宋体",
        latin="Times New Roman",
        size=12,
    )
    quote_date = data["quote_date"]
    _replace_paragraph(
        original_paragraphs[14],
        f"日期：{quote_date.year}年{quote_date.month}月{quote_date.day}日",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
        chinese="宋体",
        latin="Times New Roman",
        size=12,
    )

    for index in (2, 4, 5, 6, 7, 8, 9, 10, 11, 12):
        _remove_element(original_paragraphs[index]._element)
    for paragraph in original_paragraphs[15:]:
        _remove_element(paragraph._element)

    output = BytesIO()
    doc.save(output)
    output = _restore_template_parts(output, template)
    safe_title = re.sub(r'[<>:"/\\|?*]+', "_", data["title"]).strip(" .") or "报价单"
    return output, f"{safe_title}.docx"
