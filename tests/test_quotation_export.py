from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document


# All names and monetary values in this module are synthetic regression fixtures.


def quotation_result(client):
    response = client.post(
        "/api/quotations/calculate",
        json={
            "upper_limit": 47500,
            "vat_rate": 3,
            "surcharge_rate": 12,
            "management_rate": 16,
            "category_count": 8,
            "rounding_unit": 100,
            "vat_mode": "inclusive",
        },
    )
    assert response.status_code == 200, response.get_json()
    result = response.get_json()
    for item in result["items"]:
        item["children"] = []
    result["items"][0]["name"] = "页面策划与交互设计"
    result["items"][0]["children"] = [
        {"name": "交互原型设计", "description": "梳理页面流程与交互状态", "amount": 2400},
        {"name": "视觉规范设计", "description": "定义颜色、字体与组件规范", "amount": 2400},
    ]
    result["items"][1]["description"] = "完成第二个网页的 HTML 开发"
    return result


def export_payload(client):
    return {
        "title": "示例客户 HTML 互动教学网页报价单",
        "client_name": "示例客户",
        "project_manager": "测试负责人",
        "quote_date": "2026-08-10",
        "include_descriptions": False,
        "quotation": quotation_result(client),
    }


def test_quotation_export_downloads_template_based_docx(client, app):
    template = Path(app.config["QUOTATION_TEMPLATE"])
    before_hash = sha256(template.read_bytes()).hexdigest()

    response = client.post("/api/quotations/export", json=export_payload(client))

    assert response.status_code == 200, response.get_json()
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert ".docx" in response.headers["Content-Disposition"]
    assert response.data.startswith(b"PK")
    assert sha256(template.read_bytes()).hexdigest() == before_hash

    preserve_parts = {
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/core.xml",
        "word/_rels/document.xml.rels",
        "word/settings.xml",
        "word/styles.xml",
        "word/fontTable.xml",
        "word/theme/theme1.xml",
    }
    with ZipFile(template) as source_package, ZipFile(BytesIO(response.data)) as exported_package:
        assert exported_package.testzip() is None
        for part in preserve_parts:
            assert exported_package.read(part) == source_package.read(part), part
        assert not any(name.startswith("word/media/") for name in source_package.namelist())
        assert not any(name.startswith("word/media/") for name in exported_package.namelist())

    document = Document(BytesIO(response.data))
    assert document.paragraphs[0].text == "示例客户 HTML 互动教学网页报价单"
    assert "根据示例客户的需求" in document.paragraphs[1].text
    assert len(document.tables) == 2

    basic_table = document.tables[0]
    assert basic_table.cell(0, 0).text == "基础开发费用"
    assert len(basic_table.rows[1]._tr.tc_lst) == 3
    assert basic_table.cell(1, 0).text == "板块"
    assert basic_table.cell(1, 1).text == "项目"
    assert basic_table.cell(1, 3).text == "价格"
    assert basic_table.cell(2, 0).text == "页面策划与交互设计"
    assert len(basic_table.rows[2]._tr.tc_lst) == 3
    assert basic_table.cell(2, 1).text == "交互原型设计"
    assert basic_table.cell(2, 3).text == "2400"
    assert basic_table.cell(3, 1).text == "视觉规范设计"
    assert basic_table.cell(3, 3).text == "2400"
    assert sum(int(row.cells[3].text) for row in basic_table.rows[2:]) == 38300

    fees_table = document.tables[1]
    assert fees_table.cell(1, 2).text == "1382"
    assert fees_table.cell(2, 2).text == "166"
    assert fees_table.cell(3, 2).text == "7590"
    assert fees_table.cell(4, 2).text == "47438"

    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "项目负责人：测试负责人" in body_text
    assert "日期：2026年8月10日" in body_text
    assert "@" not in body_text
    assert "地址" not in body_text
    assert "电话" not in body_text
    assert len(document.inline_shapes) == 0

    basic_text = "\n".join(cell.text for row in basic_table.rows for cell in row.cells)
    assert "说明" not in basic_text
    assert "梳理页面流程与交互状态" not in basic_text


def test_quotation_export_includes_description_column_when_enabled(client):
    payload = export_payload(client)
    payload["include_descriptions"] = True

    response = client.post("/api/quotations/export", json=payload)

    assert response.status_code == 200, response.get_json()
    document = Document(BytesIO(response.data))
    basic_table = document.tables[0]
    assert len(basic_table.rows[1]._tr.tc_lst) == 4
    assert basic_table.cell(1, 2).text == "说明"
    assert len(basic_table.rows[2]._tr.tc_lst) == 4
    assert basic_table.cell(2, 2).text == "梳理页面流程与交互状态"
    assert basic_table.cell(3, 2).text == "定义颜色、字体与组件规范"
    assert "完成第二个网页的 HTML 开发" in basic_table.cell(4, 2).text


def test_quotation_export_rejects_inconsistent_amounts(client):
    payload = export_payload(client)
    payload["quotation"]["items"][0]["children"][0]["amount"] = 2300

    response = client.post("/api/quotations/export", json=payload)

    assert response.status_code == 400
    assert response.get_json()["error"] == "invalid_quotation_export"
    assert "子项目金额合计不一致" in response.get_json()["message"]


def test_quotation_export_rejects_non_boolean_description_setting(client):
    payload = export_payload(client)
    payload["include_descriptions"] = "false"

    response = client.post("/api/quotations/export", json=payload)

    assert response.status_code == 400
    assert response.get_json()["error"] == "invalid_quotation_export"
    assert "说明列设置格式无效" in response.get_json()["message"]
